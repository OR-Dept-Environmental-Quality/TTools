
import csv
from pathlib import Path
from openpyxl import Workbook, load_workbook

# This function was pulled from heat source v 9.1.0 at commit 89efc28
# It combines read_csv_to_list(...) and read_xlsx_to_list(...)
# Also changed to allow strings for the input path: path = Path(path)
def read_to_list(path, skiprows=0, skipcols=0, sheetname=None):
    """Read a CSV file or Excel worksheet into a list of lists indexed by row number.

    CSV values remain strings. Excel values retain their cell types, and
    formulas return their saved results. Headers are included unless skipped.
    """
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            for _ in range(skiprows):
                next(reader)
            return [row[skipcols:] for row in reader if row]

    wb = load_workbook(filename=path, data_only=True)
    ws = wb.worksheets[0] if sheetname is None else wb[sheetname]

    data = []
    for r, row in enumerate(ws.iter_rows(values_only=True)):
        if r < skiprows or row is None:
            continue
        values = list(row)[skipcols:]
        if values:
            data.append(values)
    wb.close()
    return data

# This function was copied from heat source v 9.1.0 at commit 89efc28
# Changed to allow strings for the input path: path = Path(path), and
# closes excel workbook using `finally: > wb.close`
def write_input(path, headers, rows, sheetname, csv_mode):
    """
    Write the input files to csv or xlsx.

    Writes the header row first, then writes all data rows in order.
    """
    try:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
    except PermissionError as exc:
        msg = f"Cannot create folder '{path.parent}'. Check write permissions."
        raise PermissionError(msg) from exc

    if csv_mode:
        try:
            with path.open("w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(headers)
                for row in rows:
                    writer.writerow(["" if v is None else v for v in row])
        except PermissionError as exc:
            msg = f"Cannot write file '{path}'. Check write permissions for this folder."
            raise PermissionError(msg) from exc
        return

    wb = Workbook()
    ws = wb.active
    ws.title = sheetname
    ws.append(headers)
    for row in rows:
        ws.append(row)
    try:
        wb.save(path)
    except PermissionError as exc:
        msg = f"Cannot write file '{path}'. Check write permissions for this folder."
        raise PermissionError(msg) from exc
    finally:
        wb.close()
